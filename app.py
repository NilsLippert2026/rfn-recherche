"""
RFN Sponsoring Cockpit
Recherche · Pipeline · Sponsoren · Leistungen & Preise · Analysen
Start lokal:  streamlit run app.py
"""
import io, json, os, time, uuid, threading, datetime as dt
import pandas as pd
import altair as alt
import streamlit as st
import recherche_pro as rp
import storage

st.set_page_config(page_title="RFN Sponsoring Cockpit", page_icon="🏟️", layout="wide",
                   initial_sidebar_state="expanded")

# ══════════════════════════════ Design ══════════════════════════════
G, Y, INK, MUTED, LINE = "#003D30", "#D6FF29", "#14231E", "#6B7A72", "#E6E9E2"
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
html, body, .stApp, button, input, textarea, select, [data-testid="stMarkdownContainer"] { font-family: 'Inter', system-ui, sans-serif !important; }
.stApp { background: #F4F5F0; }
#MainMenu, footer, [data-testid="stToolbar"] { visibility: hidden; }
header[data-testid="stHeader"] { background: transparent; }
.block-container { padding-top: 1.4rem; padding-bottom: 3rem; max-width: 1440px; }
section[data-testid="stSidebar"] { background: #FFFFFF; border-right: 1px solid #E6E9E2; }
h1, h2, h3, h4 { color: #0B2E26; letter-spacing: -.02em; }
.brand { display:flex; align-items:center; gap:11px; padding: 2px 0 16px; }
.brand .logo { width:40px; height:40px; border-radius:12px; background:#003D30; color:#D6FF29; display:flex; align-items:center; justify-content:center; font-weight:800; font-size:14px; letter-spacing:-.02em; }
.brand .t { font-weight:800; color:#0B2E26; font-size:16px; line-height:1.15; }
.brand .s { font-size:11px; color:#6B7A72; }
.page-h h1 { font-size:30px; font-weight:800; margin:0; }
.page-h p { color:#6B7A72; margin:4px 0 18px; font-size:14px; }
.kpi { background:#fff; border:1px solid #E6E9E2; border-radius:16px; padding:16px 18px; box-shadow:0 1px 2px rgba(16,24,40,.04); min-height:104px; }
.kpi .l { font-size:11px; font-weight:700; color:#6B7A72; text-transform:uppercase; letter-spacing:.07em; }
.kpi .v { font-size:27px; font-weight:800; color:#0B2E26; margin-top:6px; letter-spacing:-.02em; line-height:1.1; }
.kpi .s { font-size:12px; color:#8A968F; margin-top:4px; }
.kpi.accent { background: linear-gradient(135deg,#003D30 0%,#00553F 100%); border-color:#003D30; }
.kpi.accent .l { color: rgba(255,255,255,.65); } .kpi.accent .v { color:#D6FF29; } .kpi.accent .s { color: rgba(255,255,255,.6); }
.sec { font-weight:700; color:#0B2E26; font-size:15px; margin: 2px 0 6px; }
.muted { color:#6B7A72; font-size:13px; }
.warn { background:#FFF7E6; border:1px solid #F3D9A4; color:#7A5310; border-radius:12px; padding:10px 14px; font-size:13px; margin-bottom:14px; }
.ok { background:#EAF6EF; border:1px solid #BFE3CC; color:#1D6B3F; border-radius:12px; padding:10px 14px; font-size:13px; }
.stButton>button, .stDownloadButton>button { border-radius:10px; font-weight:600; }
.stButton>button[kind="primary"], .stDownloadButton>button[kind="primary"] { background:#003D30; color:#D6FF29; border:0; }
.stButton>button[kind="primary"]:hover, .stDownloadButton>button[kind="primary"]:hover { background:#00553F; color:#D6FF29; }
div[data-testid="stDataFrame"], div[data-testid="stDataEditor"] { border-radius:12px; overflow:hidden; }
.stTabs [data-baseweb="tab-list"] { gap: 6px; }
.stTabs [data-baseweb="tab"] { font-weight:600; }
div[data-testid="stVerticalBlockBorderWrapper"] { background:#fff; border-radius:16px !important; }
</style>
""", unsafe_allow_html=True)

# ══════════════════════════════ Konstanten ══════════════════════════════
PAGES = ["📊  Dashboard", "🎯  Leads & Pipeline", "🤝  Sponsoren", "💶  Leistungen & Preise", "🔍  Recherchen", "⚙️  Einstellungen"]
STATUS = ["Neu", "Anruf nötig", "Kontaktiert", "Follow-Up", "Angebot", "Gewonnen", "Verloren"]
OPEN = ["Neu", "Anruf nötig", "Kontaktiert", "Follow-Up", "Angebot"]
STATUS_COLORS = {"Neu": "#9AA79F", "Anruf nötig": "#E0A030", "Kontaktiert": "#3B82C4", "Follow-Up": "#7B61C9",
                 "Angebot": "#0F8C8C", "Gewonnen": "#1F9D55", "Verloren": "#D64545"}
SIZES = ["Kleinstunternehmen", "Kleinunternehmen", "Mittelstand", "Großunternehmen", "Filiale einer Kette"]
INFOS = ["E-Mail", "Telefon", "Ansprechpartner", "Mitarbeiterzahl", "Kurzbeschreibung",
         "Rechtsform & Handelsregister", "Gründungsjahr", "Google-Bewertung"]
WEB_INFOS = {"E-Mail", "Ansprechpartner", "Mitarbeiterzahl", "Kurzbeschreibung", "Rechtsform & Handelsregister", "Gründungsjahr"}
EINHEITEN = ["pro Saison", "pro Jahr", "pro Monat", "pro Spiel", "einmalig"]
LOST_REASONS = ["Kein Budget", "Kein Interesse", "Bereits anderweitig gebunden", "Nicht erreichbar",
                "Zeitpunkt passt nicht", "Sonstiges"]
DEFAULT_CATALOG = [
    {"name": "Bandenwerbung Hauptplatz", "kategorie": "Stadion", "preis": 450.0, "einheit": "pro Saison"},
    {"name": "Trikotwerbung Brust", "kategorie": "Trikot", "preis": 2500.0, "einheit": "pro Saison"},
    {"name": "Trikotwerbung Ärmel", "kategorie": "Trikot", "preis": 800.0, "einheit": "pro Saison"},
    {"name": "Trainingsanzug / Aufwärmshirt", "kategorie": "Trikot", "preis": 600.0, "einheit": "pro Saison"},
    {"name": "Logo auf Website & App", "kategorie": "Digital", "preis": 300.0, "einheit": "pro Saison"},
    {"name": "VereinsTV-Infoscreen Spot", "kategorie": "Digital", "preis": 400.0, "einheit": "pro Saison"},
    {"name": "Social-Media-Paket (6 Posts)", "kategorie": "Digital", "preis": 250.0, "einheit": "pro Saison"},
    {"name": "Stadionheft-Anzeige ½ Seite", "kategorie": "Print", "preis": 200.0, "einheit": "pro Saison"},
    {"name": "Spieltagspartner Heimspiel", "kategorie": "Event", "preis": 150.0, "einheit": "pro Spiel"},
    {"name": "Ballspende", "kategorie": "Event", "preis": 120.0, "einheit": "einmalig"},
]
DEFAULT_SETTINGS = {"share": 30, "provision": 40, "team": ["Nils Lippert", "Marcel", "Björn"], "google_max": 1000}


# ══════════════════════════════ Helfer ══════════════════════════════
def secret(name, default=""):
    try:
        return st.secrets.get(name, default)
    except Exception:
        return default

for _k in ("SUPABASE_URL", "SUPABASE_KEY"):
    _v = secret(_k, "")
    if _v:
        os.environ[_k] = str(_v)

def new_id(): return uuid.uuid4().hex[:10]
def today(): return dt.date.today().isoformat()
def now(): return dt.datetime.now().strftime("%d.%m.%Y %H:%M")
def eur(x):
    try: return f"{float(x):,.0f} €".replace(",", ".")
    except Exception: return "0 €"

def page_header(title, sub=""):
    st.markdown(f'<div class="page-h"><h1>{title}</h1><p>{sub}</p></div>', unsafe_allow_html=True)

def kpi(col, label, value, sub="", accent=False):
    col.markdown(f'<div class="kpi{" accent" if accent else ""}"><div class="l">{label}</div>'
                 f'<div class="v">{value}</div><div class="s">{sub}</div></div>', unsafe_allow_html=True)

def chart_style(ch):
    return (ch.configure_view(strokeWidth=0)
              .configure_axis(labelColor=MUTED, titleColor=MUTED, domainColor=LINE, tickColor=LINE,
                              gridColor="#EEF0EA", labelFont="Inter", titleFont="Inter", labelFontSize=12)
              .configure_legend(labelFont="Inter", labelColor=INK, titleFont="Inter"))

def to_excel(sheets: dict) -> bytes:
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            sn = name[:31]
            df.to_excel(xw, sheet_name=sn, index=False)
            ws = xw.sheets[sn]
            for c in ws[1]:
                c.fill = PatternFill("solid", fgColor="003D30")
                c.font = Font(bold=True, color="D6FF29")
                c.alignment = Alignment(vertical="center")
            for i, col in enumerate(df.columns, 1):
                vals = [len(str(v)) for v in df[col].head(300).tolist()]
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(str(col)) + 2, *vals))
            ws.freeze_panes = "A2"
            if len(df): ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


# ══════════════════════════════ Daten ══════════════════════════════
def DB():
    if "db" not in st.session_state:
        try:
            st.session_state.db = storage.load_all()
            st.session_state.db_error = ""
        except Exception as e:
            st.session_state.db = {}
            st.session_state.db_error = str(e)
    return st.session_state.db

def put(key, value):
    DB()[key] = value
    try:
        storage.save(key, value)
    except Exception as e:
        st.toast(f"Speichern fehlgeschlagen: {e}", icon="⚠️")

def clubs(): return DB().get("clubs", [])
def club_by_id(cid): return next((c for c in clubs() if c["id"] == cid), None)
def leads(cid): return DB().get(f"leads:{cid}", [])
def runs(): return DB().get("runs", [])
def catalog_template(): return DB().get("catalog", DEFAULT_CATALOG)
def settings():
    s = dict(DEFAULT_SETTINGS); s.update(DB().get("settings", {})); return s

def all_leads(cid="__all__"):
    out = []
    for c in clubs():
        if cid != "__all__" and c["id"] != cid:
            continue
        for l in leads(c["id"]):
            d = dict(l); d["_cid"] = c["id"]; d["_verein"] = c["name"]; out.append(d)
    return out

def split(w):
    s = settings(); w = float(w or 0)
    rfn = w * s["share"] / 100; prov = rfn * s["provision"] / 100
    return {"verein": w - rfn, "rfn_brutto": rfn, "provision": prov, "rfn_netto": rfn - prov}

def mk_lead(r):
    return {
        "id": new_id(), "firma": r["firma"], "branche": r.get("branche", ""), "typ": r.get("typ", ""),
        "adresse": r.get("adresse", ""), "dist_km": round(r["dist_km"], 1) if r.get("dist_km") is not None else None,
        "telefon": r.get("telefon", ""), "email": r.get("email", ""), "website": r.get("website", ""),
        "ansprechpartner": r.get("ansprechpartner", ""), "beschreibung": r.get("beschreibung", ""),
        "groesse": r.get("groesse", ""), "groesse_basis": r.get("groesse_basis", ""),
        "mitarbeiter": r.get("mitarbeiter"), "rechtsform": r.get("rechtsform", ""), "hrb": r.get("hrb", ""),
        "gruendung": r.get("gruendung"), "rating": r.get("rating"), "reviews": r.get("reviews"),
        "score": r.get("score", 0), "quelle": r.get("quelle", ""),
        "kontakt_status": r.get("kontakt_status") or ("E-Mail vorhanden" if r.get("email") else "nur Telefon" if r.get("telefon") else "kein Kontakt"),
        "status": "Neu" if r.get("email") else "Anruf nötig", "status_seit": today(),
        "vertriebler": "", "notiz": "", "leistungen": [], "wert": 0.0, "verlustgrund": "",
        "verlauf": [{"datum": now(), "text": "Per Recherche gefunden"}], "erstellt": today(),
    }

def new_club(name, adresse=""):
    c = {"id": new_id(), "name": name, "adresse": adresse, "erstellt": today(),
         "leistungen": [dict(s, id=new_id(), aktiv=True) for s in catalog_template()]}
    put("clubs", clubs() + [c])
    return c

def set_status(lead, status):
    if lead.get("status") != status:
        lead["verlauf"] = lead.get("verlauf", []) + [{"datum": now(), "text": f"Status: {lead.get('status')} → {status}"}]
        lead["status"] = status
        lead["status_seit"] = today()


# ══════════════════════════════ Zugang ══════════════════════════════
BRAND = '<div class="brand"><div class="logo">RFN</div><div><div class="t">Sponsoring Cockpit</div><div class="s">Regionalfußball.net</div></div></div>'
pw_required = secret("APP_PASSWORD", "")
if pw_required and not st.session_state.get("auth"):
    _, mid, _ = st.columns([1, 1.1, 1])
    with mid:
        st.markdown("<div style='height:12vh'></div>", unsafe_allow_html=True)
        st.markdown(BRAND, unsafe_allow_html=True)
        with st.container(border=True):
            st.markdown("#### Anmelden")
            pw = st.text_input("Passwort", type="password", label_visibility="collapsed", placeholder="Passwort")
            if st.button("Anmelden", type="primary", width="stretch"):
                if pw == pw_required:
                    st.session_state.auth = True; st.rerun()
                else:
                    st.error("Falsches Passwort.")
    st.stop()


# ══════════════════════════════ Recherche-Dialog ══════════════════════════════
@st.dialog("Neue Recherche", width="large")
def research_dialog():
    C = clubs()
    opts = [c["id"] for c in C] + ["__new__"]
    cur = st.session_state.get("club", "__all__")
    idx = opts.index(cur) if cur in opts else (0 if C else len(opts) - 1)

    st.markdown('<div class="sec">1 · Verein & Standort</div>', unsafe_allow_html=True)
    a, b = st.columns(2)
    sel = a.selectbox("Verein", opts, index=idx,
                      format_func=lambda x: "＋ Neuen Verein anlegen" if x == "__new__" else club_by_id(x)["name"])
    name = b.text_input("Name des neuen Vereins", placeholder="z. B. FC Bad Pyrmont Hagen") if sel == "__new__" else None
    default_addr = "" if sel == "__new__" else (club_by_id(sel) or {}).get("adresse", "")
    adresse = st.text_input("Standort (Straße Nr, PLZ Ort)", value=default_addr, placeholder="Parkstraße 7, 31812 Bad Pyrmont")
    radius = st.slider("Umkreis (km)", 1, 50, 15)

    st.markdown('<div class="sec">2 · Zielgruppe</div>', unsafe_allow_html=True)
    branchen = st.multiselect("Branchen", list(rp.BRANCHEN), default=list(rp.BRANCHEN),
                              format_func=lambda k: rp.BRANCHEN[k][0])
    groessen = st.multiselect("Unternehmensgröße", SIZES, default=SIZES,
                              help="EU-KMU: Kleinst < 10 · Klein < 50 · Mittelstand < 250 · Groß ≥ 250 Mitarbeiter. "
                                   "Ketten-Filialen sind standardmäßig ausgeschlossen (Entscheider sitzt nicht vor Ort).")

    st.markdown('<div class="sec">3 · Benötigte Informationen</div>', unsafe_allow_html=True)
    infos = st.multiselect("Welche Infos sollen recherchiert werden?", INFOS,
                           default=["E-Mail", "Telefon", "Ansprechpartner", "Mitarbeiterzahl", "Kurzbeschreibung"],
                           help="Telefon und Bewertung kommen aus der Kartenquelle, alles andere aus der Analyse der Firmen-Websites.")
    c, d = st.columns(2)
    max_n = c.number_input("Maximale Anzahl Unternehmen (0 = alle)", 0, 5000, 0, step=50,
                           help="Bei Begrenzung werden die nächstgelegenen passenden Unternehmen übernommen.")
    key = secret("GOOGLE_API_KEY", "")
    with d:
        st.markdown("<div style='height:28px'></div>", unsafe_allow_html=True)
        st.markdown(f"Quelle: **{'Google Places' if key else 'OpenStreetMap'}**"
                    + ("" if key else " · <span class='muted'>Google-Key in den Secrets aktiviert mehr Treffer</span>"),
                    unsafe_allow_html=True)

    if st.button("🚀  Recherche starten", type="primary", width="stretch"):
        if sel == "__new__" and not (name or "").strip():
            st.error("Bitte einen Vereinsnamen eingeben."); return
        if not adresse.strip():
            st.error("Bitte einen Standort eingeben."); return
        if not branchen:
            st.error("Bitte mindestens eine Branche wählen."); return
        if not groessen:
            st.error("Bitte mindestens eine Unternehmensgröße wählen."); return
        st.session_state.pending_start = dict(club=sel, name=(name or "").strip(), adresse=adresse.strip(), radius=radius,
                                              branchen=branchen, groessen=groessen, infos=infos, max_n=int(max_n),
                                              max_calls=int(settings().get("google_max", 1000)))
        st.rerun()


# ══════════════════════════════ Hintergrund-Recherche ══════════════════════════════
@st.cache_resource
def _registry():
    """Prozessweites Register: Jobs laufen unabhängig von Browser-Sitzungen weiter."""
    return {"tl": threading.local(), "jobs": {}, "lock": threading.Lock()}
REG = _registry()

def _dispatch_log(m):
    job = getattr(REG["tl"], "job", None)
    if job is not None:
        job["logs"].append(m)
        if len(job["logs"]) > 300:
            del job["logs"][:100]

def _dispatch_prog(i, n, t):
    job = getattr(REG["tl"], "job", None)
    if job is not None:
        job["i"], job["n"], job["text"] = i, n, t

rp.LOG, rp.PROGRESS = _dispatch_log, _dispatch_prog

def _merge_into(cid, rows):
    """Ergebnisse in die Vereinsdaten übernehmen (idempotent – darf mehrfach laufen)."""
    existing = storage.load(f"leads:{cid}", []) or []
    index = {(l["firma"].lower(), (l.get("adresse") or "").lower()): l for l in existing}
    new = []
    for r in rows:
        k = (r["firma"].lower(), (r.get("adresse") or "").lower())
        l = index.get(k)
        if l is not None:
            for f in ("email", "telefon", "ansprechpartner", "beschreibung", "mitarbeiter", "website",
                      "rechtsform", "hrb", "gruendung", "groesse", "groesse_basis", "score"):
                if r.get(f) and (not l.get(f) or f in ("groesse", "groesse_basis", "score")):
                    l[f] = r[f]
            l["kontakt_status"] = "E-Mail vorhanden" if l.get("email") else "nur Telefon" if l.get("telefon") else "kein Kontakt"
            if l.get("email") and l.get("status") == "Anruf nötig":
                l["status"] = "Neu"
        else:
            nl = mk_lead(r); index[k] = nl; new.append(nl)
    storage.save(f"leads:{cid}", existing + new)
    return len(new)

def _worker(job, p, cid, cname, key):
    REG["tl"].job = job
    last = {"t": 0.0}
    def checkpoint(rows, phase):
        if phase == "suche" or time.time() - last["t"] > 45:
            job["added"] += _merge_into(cid, rows)
            job["saved"] = len(rows); last["t"] = time.time()
            _dispatch_log(f"💾 Zwischenstand gespeichert ({len(rows)} Unternehmen)")
    try:
        _, rows = rp.run(cname, p["adresse"], p["radius"], p["branchen"], key,
                         scrape=bool(WEB_INFOS & set(p["infos"])), write=False,
                         max_n=p["max_n"] or None, size_filter=set(p["groessen"]),
                         max_google_calls=p.get("max_calls") or None, checkpoint=checkpoint, coords=p.get("coords"))
        job["added"] += _merge_into(cid, rows)
        s = dict(getattr(rp, "LAST_STATS", {}) or {})
        if s.get("google_fehler"):
            quelle = f"OSM {s.get('osm', 0)} (Google gesperrt)"
        elif key:
            quelle = f"Google {s.get('google', 0)} + OSM {s.get('osm', 0)}" + (" (OSM-Ausfall)" if s.get("osm_fehler") else "")
        else:
            quelle = f"OSM {s.get('osm', 0)}"
        if s.get("google_fehler"):
            job["warnung"] = s["google_fehler"]
        run_rec = {"id": job["id"], "datum": job["start_txt"], "verein": cname, "club_id": cid, "adresse": p["adresse"],
                   "radius": p["radius"], "branchen": ", ".join(rp.BRANCHEN[b][0] for b in p["branchen"]),
                   "groessen": ", ".join(p["groessen"]), "max": p["max_n"] or "alle", "quelle": quelle,
                   "google_abfragen": s.get("google_calls", 0), "gefunden": len(rows), "neu": job["added"],
                   "mit_email": sum(1 for r in rows if r.get("email"))}
        storage.save("runs", [run_rec] + (storage.load("runs", []) or [])[:199])
        cl = storage.load("clubs", []) or []
        co = s.get("coords")
        for x in cl:
            if x["id"] == cid:
                x["adresse"] = p["adresse"]
                if co:
                    x["geo"] = {"adresse": p["adresse"].strip().lower(), "lat": co[0], "lng": co[1]}
        storage.save("clubs", cl)
        job.update(status="done", found=len(rows), mails=run_rec["mit_email"], quelle=quelle)
    except SystemExit as e:
        job.update(status="error", error=str(e).strip())
    except Exception as e:
        job.update(status="error", error=f"{type(e).__name__}: {e}")
    finally:
        job["ende"] = time.time()
        REG["tl"].job = None

def start_job(p):
    club = new_club(p["name"], p["adresse"]) if p["club"] == "__new__" else club_by_id(p["club"])
    if club is None:
        st.error("Verein nicht gefunden."); return
    geo = club.get("geo") or {}
    if geo.get("adresse") == p["adresse"].strip().lower() and geo.get("lat"):
        p["coords"] = (geo["lat"], geo["lng"])
    for j in REG["jobs"].values():
        if j["club_id"] == club["id"] and j["status"] == "running":
            st.warning("Für diesen Verein läuft bereits eine Recherche."); return
    job = {"id": new_id(), "club_id": club["id"], "verein": club["name"], "adresse": p["adresse"], "radius": p["radius"],
           "status": "running", "i": 0, "n": 1, "text": "Starte …", "logs": [], "added": 0, "saved": 0,
           "start": time.time(), "start_txt": now(), "ende": None}
    with REG["lock"]:
        REG["jobs"][job["id"]] = job
    threading.Thread(target=_worker, args=(job, p, club["id"], club["name"], secret("GOOGLE_API_KEY", "")),
                     daemon=True).start()
    st.session_state.setdefault("my_jobs", set()).add(job["id"])
    st.session_state.goto_club = club["id"]

def _fmt_dur(sec):
    sec = int(sec); return f"{sec // 60} min {sec % 60:02d} s" if sec >= 60 else f"{sec} s"

@st.fragment(run_every=2)
def job_panel():
    mine = st.session_state.setdefault("my_jobs", set())
    now_t = time.time()
    shown = [j for j in REG["jobs"].values()
             if j["status"] == "running" or (not j.get("shown") and now_t - (j["ende"] or now_t) < 86400)]
    for j in sorted(shown, key=lambda x: x["start"]):
        with st.container(border=True):
            if j["status"] == "running":
                st.markdown(f"**🔄 Recherche läuft im Hintergrund** · {j['verein']} · {j['radius']} km um {j['adresse']}")
                st.progress(min(1.0, j["i"] / max(j["n"], 1)), text=f"{j['i']}/{j['n']} · {j['text'][:80]}")
                a, b = st.columns([3, 1])
                a.caption(f"Laufzeit {_fmt_dur(now_t - j['start'])} · gesichert: {j['saved']} Unternehmen · "
                          "Du kannst weiterarbeiten oder den Tab schließen – die Recherche läuft weiter.")
                if b.button("Zwischenstand laden", key=f"zw_{j['id']}"):
                    st.session_state.pop("db", None); st.rerun()
                with st.expander("Protokoll"):
                    st.markdown("<br>".join(f"<span style='font-size:12px'>{x}</span>" for x in j["logs"][-15:]),
                                unsafe_allow_html=True)
            else:
                j["shown"] = True
                st.session_state.pop("db", None)
                if j["status"] == "done":
                    st.session_state.flash = (f"✓ Recherche {j['verein']} abgeschlossen in {_fmt_dur(j['ende'] - j['start'])}: "
                                              f"{j.get('found', 0)} Unternehmen · {j.get('mails', 0)} mit E-Mail · "
                                              f"{j['added']} neu übernommen · Quelle: {j.get('quelle', '')}")
                    if j.get("warnung"):
                        st.session_state.flash_error = ("Google hat den Zugriff verweigert – die Recherche lief nur mit OpenStreetMap. "
                                                        f"Grund laut Google: {j['warnung']}")
                    if j["id"] in mine:
                        st.session_state.goto_club = j["club_id"]; st.session_state.goto_nav = PAGES[1]
                else:
                    st.session_state.flash_error = (f"Recherche {j['verein']} abgebrochen: {j.get('error', '').strip()} – "
                                                    f"bereits gesicherte {j['saved']} Unternehmen bleiben erhalten.\n\n"
                                                    "Protokoll: " + " · ".join(j["logs"][-6:]))
                st.rerun()


# ══════════════════════════════ Seiten ══════════════════════════════
def empty_state(text="Noch keine Daten – starte die erste Recherche."):
    with st.container(border=True):
        st.markdown(f"### 👋 Los geht's\n{text}")
        if st.button("🔍  Erste Recherche starten", type="primary"):
            research_dialog()


# ---------- Dashboard ----------
def page_dashboard():
    cid = st.session_state.club
    rows = all_leads(cid)
    scope = "alle Vereine" if cid == "__all__" else club_by_id(cid)["name"]
    page_header("Dashboard", f"Überblick · {scope}")
    if not rows:
        empty_state(); return
    df = pd.DataFrame(rows)
    won, lost = df[df.status == "Gewonnen"], df[df.status == "Verloren"]
    open_ = df[df.status.isin(OPEN)]
    vol = float(won["wert"].sum()); sp = split(vol)
    rate = len(won) / (len(won) + len(lost)) if (len(won) + len(lost)) else None
    mail = int((df["email"].fillna("") != "").sum())

    c = st.columns(6)
    kpi(c[0], "Unternehmen", f"{len(df)}", f"{df['_cid'].nunique()} Verein(e)")
    kpi(c[1], "Mit E-Mail", f"{mail}", f"{mail / len(df) * 100:.0f} % der Firmen")
    kpi(c[2], "In Bearbeitung", f"{len(open_[open_.status.isin(['Kontaktiert', 'Follow-Up', 'Angebot'])])}", "Kontaktiert · Follow-Up · Angebot")
    kpi(c[3], "Gewonnen", f"{len(won)}", "Sponsoren")
    kpi(c[4], "Verloren", f"{len(lost)}", "abgesagt")
    kpi(c[5], "Abschlussquote", f"{rate * 100:.0f} %" if rate is not None else "–", "gewonnen / entschieden")
    st.write("")
    c = st.columns(4)
    kpi(c[0], "Sponsoring-Volumen", eur(vol), "gebuchte Leistungen", accent=True)
    kpi(c[1], f"an Vereine ({100 - settings()['share']} %)", eur(sp["verein"]), "Vereinsanteil")
    kpi(c[2], "RFN netto", eur(sp["rfn_netto"]), f"{settings()['share']} % abzgl. Provision")
    kpi(c[3], "Vertriebsprovision", eur(sp["provision"]), f"{settings()['provision']} % vom RFN-Anteil")
    st.write("")

    l, r = st.columns(2)
    with l.container(border=True):
        st.markdown('<div class="sec">Pipeline</div>', unsafe_allow_html=True)
        d = df.groupby("status").size().reindex(STATUS, fill_value=0).reset_index()
        d.columns = ["Status", "Anzahl"]
        ch = alt.Chart(d).mark_bar(cornerRadiusEnd=6, height=22).encode(
            x=alt.X("Anzahl:Q", title=None), y=alt.Y("Status:N", sort=STATUS, title=None),
            color=alt.Color("Status:N", scale=alt.Scale(domain=STATUS, range=[STATUS_COLORS[s] for s in STATUS]), legend=None),
            tooltip=["Status", "Anzahl"]).properties(height=270)
        st.altair_chart(chart_style(ch), width="stretch")
    with r.container(border=True):
        st.markdown('<div class="sec">Gewonnen vs. verloren pro Monat</div>', unsafe_allow_html=True)
        wl = df[df.status.isin(["Gewonnen", "Verloren"])].copy()
        if len(wl):
            wl["Monat"] = wl["status_seit"].astype(str).str[:7]
            d = wl.groupby(["Monat", "status"]).size().reset_index(name="Anzahl")
            ch = alt.Chart(d).mark_bar(cornerRadiusTopLeft=5, cornerRadiusTopRight=5).encode(
                x=alt.X("Monat:N", title=None), xOffset="status:N", y=alt.Y("Anzahl:Q", title=None),
                color=alt.Color("status:N", scale=alt.Scale(domain=["Gewonnen", "Verloren"], range=["#1F9D55", "#D64545"]),
                                legend=alt.Legend(title=None, orient="top")),
                tooltip=["Monat", "status", "Anzahl"]).properties(height=270)
            st.altair_chart(chart_style(ch), width="stretch")
        else:
            st.markdown('<p class="muted">Noch keine Entscheidungen – sobald Leads gewonnen oder verloren sind, erscheint hier der Verlauf.</p>', unsafe_allow_html=True)

    l, r = st.columns(2)
    with l.container(border=True):
        st.markdown('<div class="sec">Branchen</div>', unsafe_allow_html=True)
        b = df.copy()
        b["Phase"] = b.status.map(lambda s: s if s in ("Gewonnen", "Verloren") else "Offen")
        d = b.groupby(["branche", "Phase"]).size().reset_index(name="Anzahl")
        ch = alt.Chart(d).mark_bar(cornerRadiusEnd=4).encode(
            x=alt.X("Anzahl:Q", title=None, stack=True), y=alt.Y("branche:N", title=None, sort="-x"),
            color=alt.Color("Phase:N", scale=alt.Scale(domain=["Offen", "Gewonnen", "Verloren"], range=["#B9C4BD", "#1F9D55", "#D64545"]),
                            legend=alt.Legend(title=None, orient="top")),
            tooltip=["branche", "Phase", "Anzahl"]).properties(height=300)
        st.altair_chart(chart_style(ch), width="stretch")
    with r.container(border=True):
        st.markdown('<div class="sec">Unternehmensgröße</div>', unsafe_allow_html=True)
        d = df.groupby("groesse").size().reindex([s for s in SIZES if s in set(df.groesse)], fill_value=0).reset_index()
        d.columns = ["Größe", "Anzahl"]
        ch = alt.Chart(d).mark_arc(innerRadius=70, cornerRadius=4).encode(
            theta="Anzahl:Q",
            color=alt.Color("Größe:N", scale=alt.Scale(domain=SIZES, range=["#C9D6CE", "#7FB89A", "#2E8B62", "#003D30", "#E0A030"]),
                            legend=alt.Legend(title=None, orient="right")),
            tooltip=["Größe", "Anzahl"]).properties(height=300)
        st.altair_chart(chart_style(ch), width="stretch")

    l, r = st.columns(2)
    with l.container(border=True):
        st.markdown('<div class="sec">Top-Leistungen nach Umsatz</div>', unsafe_allow_html=True)
        items = [{"Leistung": b["name"], "Umsatz": float(b["preis"]) * float(b["menge"]), "Anzahl": float(b["menge"])}
                 for _, w in won.iterrows() for b in (w["leistungen"] or [])]
        if items:
            d = pd.DataFrame(items).groupby("Leistung", as_index=False).sum().sort_values("Umsatz", ascending=False).head(10)
            ch = alt.Chart(d).mark_bar(cornerRadiusEnd=6, color=G).encode(
                x=alt.X("Umsatz:Q", title=None), y=alt.Y("Leistung:N", sort="-x", title=None),
                tooltip=["Leistung", alt.Tooltip("Umsatz:Q", format=",.0f"), "Anzahl"]).properties(height=280)
            st.altair_chart(chart_style(ch), width="stretch")
        else:
            st.markdown('<p class="muted">Noch keine gebuchten Leistungen.</p>', unsafe_allow_html=True)
    with r.container(border=True):
        if cid == "__all__":
            st.markdown('<div class="sec">Volumen pro Verein</div>', unsafe_allow_html=True)
            d = won.groupby("_verein", as_index=False)["wert"].sum().rename(columns={"_verein": "Verein", "wert": "Volumen"})
        else:
            st.markdown('<div class="sec">Provision pro Vertriebler</div>', unsafe_allow_html=True)
            w2 = won.copy(); w2["Vertriebler"] = w2["vertriebler"].replace("", "– offen –")
            w2["Provision"] = w2["wert"].map(lambda x: split(x)["provision"])
            d = w2.groupby("Vertriebler", as_index=False)["Provision"].sum().rename(columns={"Provision": "Volumen"})
            d = d.rename(columns={"Vertriebler": "Verein"})
        if len(d) and d["Volumen"].sum() > 0:
            ch = alt.Chart(d).mark_bar(cornerRadiusEnd=6, color="#2E8B62").encode(
                x=alt.X("Volumen:Q", title=None), y=alt.Y("Verein:N", sort="-x", title=None),
                tooltip=["Verein", alt.Tooltip("Volumen:Q", format=",.0f")]).properties(height=280)
            st.altair_chart(chart_style(ch), width="stretch")
        else:
            st.markdown('<p class="muted">Noch kein Umsatz.</p>', unsafe_allow_html=True)


# ---------- Leads & Pipeline ----------
LEAD_COLS = {"firma": "Firma", "_verein": "Verein", "status": "Status", "kontakt_status": "Kontakt", "vertriebler": "Vertriebler", "wert": "Wert",
             "email": "E-Mail", "telefon": "Telefon", "ansprechpartner": "Ansprechpartner", "branche": "Branche",
             "groesse": "Größe", "dist_km": "km", "mitarbeiter": "Mitarbeiter", "beschreibung": "Beschreibung",
             "website": "Website", "adresse": "Adresse", "rechtsform": "Rechtsform", "gruendung": "Gegründet",
             "rating": "Bewertung", "score": "Score", "status_seit": "Status seit"}
EDITABLE = ["Status", "Vertriebler", "E-Mail", "Telefon", "Ansprechpartner"]

def leads_frame(rows, cols):
    df = pd.DataFrame(rows)
    df.index = df["_cid"] + "|" + df["id"]
    keep = [k for k in LEAD_COLS if LEAD_COLS[k] in cols and k in df.columns]
    out = df[keep].rename(columns=LEAD_COLS)
    for c in ("E-Mail", "Telefon", "Ansprechpartner", "Vertriebler"):
        if c in out: out[c] = out[c].fillna("")
    return out

def export_leads(rows):
    df = pd.DataFrame(rows)
    cols = [k for k in LEAD_COLS if k in df.columns] + ["groesse_basis", "hrb", "reviews", "notiz", "verlustgrund"]
    out = df[[c for c in cols if c in df.columns]].rename(columns=dict(LEAD_COLS, groesse_basis="Größe – Basis", hrb="Handelsregister",
                                                                       reviews="Bewertungen", notiz="Notiz", verlustgrund="Verlustgrund"))
    out["Gebuchte Leistungen"] = df["leistungen"].map(lambda L: ", ".join(f"{int(b['menge'])}× {b['name']}" for b in (L or [])))
    return out

def page_leads():
    cid = st.session_state.club
    rows = all_leads(cid)
    scope = "alle Vereine" if cid == "__all__" else club_by_id(cid)["name"]
    page_header("Leads & Pipeline", f"Unternehmen bearbeiten, Status pflegen, Leistungen buchen · {scope}")
    if not rows:
        empty_state(); return
    team = settings()["team"]

    with st.container(border=True):
        f = st.columns([2, 1.4, 1.4, 1.3, 1.3, 1.1])
        q = f[0].text_input("Suche", placeholder="Firma, Ort, Ansprechpartner …")
        fk = f[1].multiselect("Kontakt", ["E-Mail vorhanden", "nur Telefon", "kein Kontakt"], placeholder="alle",
                              help="„kein Kontakt“ ausblenden, um nur erreichbare Leads zu sehen.")
        fs = f[2].multiselect("Status", STATUS, placeholder="alle")
        fb = f[3].multiselect("Branche", sorted({r["branche"] for r in rows}), placeholder="alle")
        fg = f[4].multiselect("Größe", [s for s in SIZES if s in {r["groesse"] for r in rows}], placeholder="alle")
        fv = f[5].multiselect("Vertriebler", ["– offen –"] + team, placeholder="alle")
        default_cols = ["Firma", "Kontakt", "Status", "Vertriebler", "E-Mail", "Telefon", "Ansprechpartner", "Branche", "Größe",
                        "km", "Mitarbeiter", "Wert", "Score"] + (["Verein"] if cid == "__all__" else [])
        cols = st.multiselect("Spalten", list(LEAD_COLS.values()), default=[c for c in LEAD_COLS.values() if c in default_cols])

    v = rows
    if q:
        ql = q.lower()
        v = [r for r in v if ql in f"{r['firma']} {r.get('adresse', '')} {r.get('ansprechpartner', '')} {r.get('beschreibung', '')}".lower()]
    if fk: v = [r for r in v if (r.get("kontakt_status") or "kein Kontakt") in fk]
    if fs: v = [r for r in v if r["status"] in fs]
    if fb: v = [r for r in v if r["branche"] in fb]
    if fg: v = [r for r in v if r["groesse"] in fg]
    if fv: v = [r for r in v if (r.get("vertriebler") or "– offen –") in fv]
    v = sorted(v, key=lambda r: (STATUS.index(r["status"]) if r["status"] in STATUS else 9, -(r.get("score") or 0)))

    if not v:
        st.info("Keine Treffer für diese Filter."); return
    if "Firma" not in cols: cols = ["Firma"] + cols
    df = leads_frame(v, cols)
    ekey = "ed_" + str(abs(hash(tuple(df.index) + tuple(df.columns))))[:12]
    edited = st.data_editor(
        df, key=ekey, hide_index=True, width="stretch", height=min(620, 42 + 35 * len(df)),
        disabled=[c for c in df.columns if c not in EDITABLE],
        column_config={
            "Status": st.column_config.SelectboxColumn("Status", options=STATUS, required=True, width="small"),
            "Vertriebler": st.column_config.SelectboxColumn("Vertriebler", options=[""] + team, width="small"),
            "Wert": st.column_config.NumberColumn("Wert", format="%.0f €", width="small"),
            "Score": st.column_config.ProgressColumn("Score", min_value=0, max_value=100, format="%d", width="small"),
            "Website": st.column_config.LinkColumn("Website"),
            "km": st.column_config.NumberColumn("km", format="%.1f", width="small"),
            "Beschreibung": st.column_config.TextColumn("Beschreibung", width="large"),
        })
    ed = [c for c in EDITABLE if c in df.columns]
    changed = (df[ed].fillna("").astype(str) != edited[ed].fillna("").astype(str)).any(axis=1) if ed else pd.Series(False, index=df.index)
    n_changed = int(changed.sum())
    a, b, c, d = st.columns([1.4, 1.4, 1.4, 3])
    if a.button(f"💾  {n_changed} Änderung(en) speichern" if n_changed else "Keine Änderungen", type="primary", disabled=not n_changed):
        touched = set()
        fmap = {"Status": "status", "Vertriebler": "vertriebler", "E-Mail": "email", "Telefon": "telefon", "Ansprechpartner": "ansprechpartner"}
        for idx in changed[changed].index:
            lcid, lid = idx.split("|")
            L = leads(lcid); lead = next(x for x in L if x["id"] == lid)
            for col in ed:
                val = edited.at[idx, col]; val = "" if pd.isna(val) else str(val)
                if col == "Status": set_status(lead, val)
                else: lead[fmap[col]] = val
            if lead.get("email") and lead["status"] == "Anruf nötig": set_status(lead, "Neu")
            touched.add(lcid)
        for lcid in touched: put(f"leads:{lcid}", leads(lcid))
        st.session_state.pop(ekey, None)
        st.toast(f"{n_changed} Änderung(en) gespeichert", icon="✅"); st.rerun()
    b.download_button("⬇️  Excel (Auswahl)", to_excel({"Leads": export_leads(v)}), file_name=f"Leads_{today()}.xlsx", width="stretch")
    c.download_button("⬇️  CSV (Auswahl)", export_leads(v).to_csv(index=False, sep=";").encode("utf-8-sig"),
                      file_name=f"Leads_{today()}.csv", width="stretch")
    d.markdown(f'<p class="muted" style="padding-top:8px">{len(v)} von {len(rows)} Unternehmen · Status, Vertriebler und Kontaktdaten direkt in der Tabelle editierbar</p>',
               unsafe_allow_html=True)

    st.write("")
    sel = st.selectbox("Unternehmen öffnen", ["—"] + list(df.index),
                       format_func=lambda i: "— Unternehmen für Details & Leistungsbuchung auswählen —" if i == "—" else
                       f"{df.at[i, 'Firma']}  ·  {next(r['status'] for r in v if r['_cid'] + '|' + r['id'] == i)}")
    if sel != "—":
        lead_detail(*sel.split("|"))


def lead_detail(cid, lid):
    club = club_by_id(cid); L = leads(cid)
    lead = next((x for x in L if x["id"] == lid), None)
    if not lead: return
    with st.container(border=True):
        h1, h2 = st.columns([3, 1])
        h1.markdown(f"### {lead['firma']}")
        h1.markdown(f'<span class="muted">{lead.get("branche", "")} · {lead.get("groesse", "")} · {club["name"]}'
                    f'{" · " + str(lead["dist_km"]) + " km" if lead.get("dist_km") is not None else ""}</span>', unsafe_allow_html=True)
        col = STATUS_COLORS.get(lead["status"], MUTED)
        h2.markdown(f"<div style='text-align:right;padding-top:14px'><span style='background:{col}1F;color:{col};padding:5px 12px;"
                    f"border-radius:999px;font-weight:700;font-size:13px'>{lead['status']}</span></div>", unsafe_allow_html=True)

        left, right = st.columns([1.1, 1])
        with left:
            st.markdown('<div class="sec">Unternehmensinfos</div>', unsafe_allow_html=True)
            info = [("📍 Adresse", lead.get("adresse")), ("📞 Telefon", lead.get("telefon")), ("✉️ E-Mail", lead.get("email")),
                    ("👤 Ansprechpartner", lead.get("ansprechpartner")), ("🌐 Website", lead.get("website")),
                    ("👥 Mitarbeiter", lead.get("mitarbeiter")), ("🏛️ Rechtsform", " · ".join(x for x in [lead.get("rechtsform"), lead.get("hrb")] if x)),
                    ("📅 Gegründet", lead.get("gruendung")),
                    ("⭐ Bewertung", f"{lead['rating']} ({lead.get('reviews')})" if lead.get("rating") else ""),
                    ("📏 Größe – Basis", lead.get("groesse_basis"))]
            st.markdown("\n".join(f"- **{k}:** {v}" for k, v in info if v not in (None, "")) or "_Keine Detaildaten_")
            if lead.get("beschreibung"):
                st.markdown(f"> {lead['beschreibung']}")
        with right:
            st.markdown('<div class="sec">Bearbeitung</div>', unsafe_allow_html=True)
            team = settings()["team"]
            status = st.selectbox("Status", STATUS, index=STATUS.index(lead["status"]) if lead["status"] in STATUS else 0, key=f"st_{lid}")
            vtr = st.selectbox("Vertriebler", [""] + team, index=([""] + team).index(lead.get("vertriebler", "")) if lead.get("vertriebler", "") in [""] + team else 0,
                               key=f"vt_{lid}", format_func=lambda x: x or "– offen –")
            grund = ""
            if status == "Verloren":
                g0 = lead.get("verlustgrund") or LOST_REASONS[0]
                grund = st.selectbox("Verlustgrund", LOST_REASONS, index=LOST_REASONS.index(g0) if g0 in LOST_REASONS else 0, key=f"vg_{lid}")
            notiz = st.text_area("Notiz", lead.get("notiz", ""), key=f"nz_{lid}", height=90)

        st.markdown('<div class="sec">Gebuchte Leistungen</div>', unsafe_allow_html=True)
        cat = [s for s in club.get("leistungen", []) if s.get("aktiv", True)]
        booked = {b["name"]: b for b in lead.get("leistungen", [])}
        names = [s["name"] for s in cat] + [n for n in booked if n not in {s["name"] for s in cat}]
        chosen = st.multiselect("Leistungen auswählen", names, default=list(booked), key=f"ms_{lid}",
                                placeholder="Leistungen aus dem Vereinskatalog wählen …")
        lines = []
        for n in chosen:
            if n in booked: lines.append(dict(booked[n]))
            else:
                s = next(s for s in cat if s["name"] == n)
                lines.append({"name": n, "menge": 1, "preis": float(s.get("preis") or 0), "einheit": s.get("einheit", "")})
        bdf = pd.DataFrame(lines, columns=["name", "menge", "preis", "einheit"])
        bed = st.data_editor(bdf, hide_index=True, width="stretch", key=f"bk_{lid}_{abs(hash(tuple(chosen)))}",
                             disabled=["name", "einheit"],
                             column_config={"name": "Leistung",
                                            "menge": st.column_config.NumberColumn("Menge", min_value=1, step=1, format="%d"),
                                            "preis": st.column_config.NumberColumn("Preis (€)", min_value=0.0, step=10.0, format="%.2f €"),
                                            "einheit": "Einheit"})
        wert = float((bed["menge"].fillna(0) * bed["preis"].fillna(0)).sum()) if len(bed) else 0.0
        sp = split(wert)
        k = st.columns(4)
        kpi(k[0], "Sponsoring-Wert", eur(wert), "Summe Leistungen", accent=True)
        kpi(k[1], "Verein", eur(sp["verein"]), f"{100 - settings()['share']} %")
        kpi(k[2], "RFN netto", eur(sp["rfn_netto"]))
        kpi(k[3], "Provision", eur(sp["provision"]), vtr or "kein Vertriebler")
        st.write("")

        b1, b2, b3, b4 = st.columns([1.3, 1.2, 1.2, 1])
        if b1.button("💾  Speichern", type="primary", key=f"sv_{lid}", width="stretch"):
            set_status(lead, status)
            lead.update(vertriebler=vtr, notiz=notiz, verlustgrund=grund if status == "Verloren" else lead.get("verlustgrund", ""),
                        leistungen=[{"name": r["name"], "menge": int(r["menge"] or 1), "preis": float(r["preis"] or 0), "einheit": r["einheit"]}
                                    for _, r in bed.iterrows()], wert=wert)
            put(f"leads:{cid}", L); st.toast("Gespeichert", icon="✅"); st.rerun()
        if b2.button("🏆  Gewonnen", key=f"gw_{lid}", width="stretch", disabled=lead["status"] == "Gewonnen"):
            set_status(lead, "Gewonnen")
            lead.update(vertriebler=vtr, notiz=notiz, wert=wert,
                        leistungen=[{"name": r["name"], "menge": int(r["menge"] or 1), "preis": float(r["preis"] or 0), "einheit": r["einheit"]}
                                    for _, r in bed.iterrows()])
            put(f"leads:{cid}", L); st.balloons(); st.rerun()
        if b3.button("✕  Verloren", key=f"vl_{lid}", width="stretch", disabled=lead["status"] == "Verloren"):
            set_status(lead, "Verloren"); lead["verlustgrund"] = lead.get("verlustgrund") or "Kein Interesse"
            put(f"leads:{cid}", L); st.rerun()
        if b4.button("🗑️", key=f"dl_{lid}", width="stretch", help="Unternehmen löschen"):
            st.session_state[f"confirm_{lid}"] = True
        if st.session_state.get(f"confirm_{lid}"):
            st.warning(f"„{lead['firma']}“ wirklich löschen?")
            y, n = st.columns(2)
            if y.button("Ja, löschen", key=f"y_{lid}"):
                put(f"leads:{cid}", [x for x in L if x["id"] != lid]); st.session_state.pop(f"confirm_{lid}"); st.rerun()
            if n.button("Abbrechen", key=f"n_{lid}"):
                st.session_state.pop(f"confirm_{lid}"); st.rerun()

        with st.expander(f"Verlauf ({len(lead.get('verlauf', []))})"):
            for e in reversed(lead.get("verlauf", [])[-30:]):
                st.markdown(f"<span class='muted'>{e['datum']}</span> · {e['text']}", unsafe_allow_html=True)
            t = st.text_input("Eintrag hinzufügen", key=f"vh_{lid}", placeholder="z. B. Telefonat mit Geschäftsführer, Angebot bis Freitag")
            if st.button("Hinzufügen", key=f"va_{lid}") and t.strip():
                lead["verlauf"] = lead.get("verlauf", []) + [{"datum": now(), "text": t.strip()}]
                put(f"leads:{cid}", L); st.rerun()


# ---------- Sponsoren ----------
def page_sponsoren():
    cid = st.session_state.club
    rows = all_leads(cid)
    scope = "alle Vereine" if cid == "__all__" else club_by_id(cid)["name"]
    page_header("Sponsoren", f"Gewonnene und verlorene Unternehmen, gebuchte Leistungen · {scope}")
    won = [r for r in rows if r["status"] == "Gewonnen"]
    lost = [r for r in rows if r["status"] == "Verloren"]
    vol = sum(float(r.get("wert") or 0) for r in won); sp = split(vol)
    c = st.columns(5)
    kpi(c[0], "Gewonnen", len(won), "aktive Sponsoren")
    kpi(c[1], "Volumen", eur(vol), "pro Jahr/Saison", accent=True)
    kpi(c[2], "Ø pro Sponsor", eur(vol / len(won)) if won else "–")
    kpi(c[3], "Verloren", len(lost))
    kpi(c[4], "Abschlussquote", f"{len(won) / (len(won) + len(lost)) * 100:.0f} %" if (won or lost) else "–")
    st.write("")

    t1, t2, t3 = st.tabs([f"🏆 Gewonnen ({len(won)})", f"✕ Verloren ({len(lost)})", "📦 Leistungsübersicht"])
    with t1:
        if not won:
            st.info("Noch keine gewonnenen Sponsoren. In „Leads & Pipeline“ ein Unternehmen öffnen, Leistungen buchen und auf „Gewonnen“ klicken.")
        else:
            df = pd.DataFrame([{
                "Firma": r["firma"], "Verein": r["_verein"],
                "Leistungen": ", ".join(f"{int(b['menge'])}× {b['name']}" for b in r.get("leistungen", [])) or "–",
                "Wert": float(r.get("wert") or 0), "Verein-Anteil": split(r.get("wert"))["verein"],
                "RFN netto": split(r.get("wert"))["rfn_netto"], "Provision": split(r.get("wert"))["provision"],
                "Vertriebler": r.get("vertriebler") or "–", "Gewonnen am": r.get("status_seit", ""),
                "Ansprechpartner": r.get("ansprechpartner", ""), "E-Mail": r.get("email", ""), "Telefon": r.get("telefon", ""),
            } for r in sorted(won, key=lambda x: -float(x.get("wert") or 0))])
            money = {k: st.column_config.NumberColumn(k, format="%.0f €") for k in ["Wert", "Verein-Anteil", "RFN netto", "Provision"]}
            st.dataframe(df, hide_index=True, width="stretch", column_config=money,
                         column_order=[c for c in df.columns if c != "Verein" or cid == "__all__"])
            st.download_button("⬇️  Sponsorenliste (Excel)", to_excel({"Sponsoren": df}), file_name=f"Sponsoren_{today()}.xlsx", type="primary")
    with t2:
        if not lost:
            st.info("Keine verlorenen Unternehmen.")
        else:
            df = pd.DataFrame([{"Firma": r["firma"], "Verein": r["_verein"], "Grund": r.get("verlustgrund") or "–",
                                "Datum": r.get("status_seit", ""), "Vertriebler": r.get("vertriebler") or "–", "Notiz": r.get("notiz", "")} for r in lost])
            a, b = st.columns([1.4, 1])
            a.dataframe(df, hide_index=True, width="stretch")
            d = df.groupby("Grund").size().reset_index(name="Anzahl")
            ch = alt.Chart(d).mark_bar(cornerRadiusEnd=6, color="#D64545").encode(
                x=alt.X("Anzahl:Q", title=None), y=alt.Y("Grund:N", sort="-x", title=None), tooltip=["Grund", "Anzahl"]).properties(height=240)
            with b.container(border=True):
                st.markdown('<div class="sec">Verlustgründe</div>', unsafe_allow_html=True)
                st.altair_chart(chart_style(ch), width="stretch")
            st.download_button("⬇️  Verlorene (Excel)", to_excel({"Verloren": df}), file_name=f"Verloren_{today()}.xlsx")
    with t3:
        items = [{"Leistung": b["name"], "Verein": r["_verein"], "Menge": float(b["menge"]), "Umsatz": float(b["preis"]) * float(b["menge"])}
                 for r in won for b in r.get("leistungen", [])]
        if not items:
            st.info("Noch keine gebuchten Leistungen.")
        else:
            d = pd.DataFrame(items)
            agg = d.groupby("Leistung").agg(Buchungen=("Menge", "sum"), Umsatz=("Umsatz", "sum"), Vereine=("Verein", "nunique")).reset_index()
            st.dataframe(agg.sort_values("Umsatz", ascending=False), hide_index=True, width="stretch",
                         column_config={"Umsatz": st.column_config.NumberColumn(format="%.0f €"), "Buchungen": st.column_config.NumberColumn(format="%d")})


# ---------- Leistungen & Preise ----------
def catalog_editor(items, key):
    df = pd.DataFrame([{"id": s.get("id", ""), "aktiv": s.get("aktiv", True), "name": s["name"], "kategorie": s.get("kategorie", ""),
                        "preis": float(s.get("preis") or 0), "einheit": s.get("einheit", "pro Saison"), "beschreibung": s.get("beschreibung", "")}
                       for s in items], columns=["id", "aktiv", "name", "kategorie", "preis", "einheit", "beschreibung"])
    ed = st.data_editor(df, key=key, num_rows="dynamic", hide_index=True, width="stretch",
                        column_config={"id": None,
                                       "aktiv": st.column_config.CheckboxColumn("Aktiv", default=True, width="small"),
                                       "name": st.column_config.TextColumn("Leistung", required=True, width="medium"),
                                       "kategorie": st.column_config.TextColumn("Kategorie", width="small"),
                                       "preis": st.column_config.NumberColumn("Preis", min_value=0.0, step=10.0, format="%.2f €", required=True),
                                       "einheit": st.column_config.SelectboxColumn("Einheit", options=EINHEITEN, default="pro Saison"),
                                       "beschreibung": st.column_config.TextColumn("Beschreibung", width="large")})
    out = []
    for _, r in ed.iterrows():
        if not str(r["name"] or "").strip():
            continue
        out.append({"id": r["id"] if isinstance(r["id"], str) and r["id"] else new_id(), "aktiv": bool(r["aktiv"]) if pd.notna(r["aktiv"]) else True,
                    "name": str(r["name"]).strip(), "kategorie": "" if pd.isna(r["kategorie"]) else str(r["kategorie"]),
                    "preis": float(r["preis"]) if pd.notna(r["preis"]) else 0.0, "einheit": r["einheit"] if pd.notna(r["einheit"]) else "pro Saison",
                    "beschreibung": "" if pd.isna(r["beschreibung"]) else str(r["beschreibung"])})
    return out

def page_leistungen():
    page_header("Leistungen & Preise", "Sponsoring-Leistungen pro Verein auswählen, ergänzen und bepreisen")
    C = clubs()
    t1, t2 = st.tabs(["🏟️ Leistungen pro Verein", "📋 Globale Vorlage"])
    with t1:
        if not C:
            empty_state("Lege zuerst über eine Recherche einen Verein an – er bekommt automatisch die Vorlage als Leistungskatalog."); 
        else:
            cur = st.session_state.club if st.session_state.club in [c["id"] for c in C] else C[0]["id"]
            cid = st.selectbox("Verein", [c["id"] for c in C], index=[c["id"] for c in C].index(cur), format_func=lambda x: club_by_id(x)["name"])
            club = club_by_id(cid)
            st.markdown('<p class="muted">Haken bei „Aktiv“ = im Verein buchbar. Neue Zeile unten hinzufügen, Zeile markieren + Entf zum Löschen. '
                        'Preise sind Listenpreise – beim Buchen pro Sponsor individuell anpassbar.</p>', unsafe_allow_html=True)
            new = catalog_editor(club.get("leistungen", []), key=f"cat_{cid}")
            a, b, _ = st.columns([1.2, 1.6, 3])
            if a.button("💾  Speichern", type="primary", key=f"savecat_{cid}"):
                for c in C:
                    if c["id"] == cid: c["leistungen"] = new
                put("clubs", C); st.session_state.pop(f"cat_{cid}", None); st.toast("Leistungskatalog gespeichert", icon="✅"); st.rerun()
            if b.button("＋ Fehlende aus Vorlage ergänzen", key=f"tpl_{cid}"):
                have = {s["name"] for s in club.get("leistungen", [])}
                add = [dict(s, id=new_id(), aktiv=True) for s in catalog_template() if s["name"] not in have]
                for c in C:
                    if c["id"] == cid: c["leistungen"] = c.get("leistungen", []) + add
                put("clubs", C); st.session_state.pop(f"cat_{cid}", None); st.toast(f"{len(add)} Leistungen ergänzt", icon="✅"); st.rerun()
    with t2:
        st.markdown('<p class="muted">Die Vorlage wird jedem neuen Verein als Startkatalog mitgegeben. Bestehende Vereine ändern sich dadurch nicht.</p>',
                    unsafe_allow_html=True)
        new = catalog_editor(catalog_template(), key="cat_tpl")
        if st.button("💾  Vorlage speichern", type="primary"):
            put("catalog", [{k: v for k, v in s.items() if k != "id"} for s in new])
            st.session_state.pop("cat_tpl", None); st.toast("Vorlage gespeichert", icon="✅"); st.rerun()


# ---------- Recherchen ----------
def page_runs():
    page_header("Recherchen", "Historie aller Suchläufe")
    if st.button("🔍  Neue Recherche", type="primary"):
        research_dialog()
    R = runs()
    if not R:
        st.info("Noch keine Recherchen."); return
    df = pd.DataFrame(R)
    if "google_abfragen" not in df: df["google_abfragen"] = None
    df = df[["datum", "verein", "adresse", "radius", "quelle", "google_abfragen", "gefunden", "neu", "mit_email", "branchen", "groessen", "max"]]
    df.columns = ["Datum", "Verein", "Standort", "km", "Quelle", "Google-Abfragen", "Gefunden", "Neu übernommen", "Mit E-Mail", "Branchen", "Größen", "Max."]
    st.dataframe(df, hide_index=True, width="stretch")


# ---------- Einstellungen ----------
def page_settings():
    page_header("Einstellungen", "Provisionsmodell, Team, Vereine, Daten")
    s = settings()
    l, r = st.columns(2)
    with l.container(border=True):
        st.markdown('<div class="sec">Provisionsmodell</div>', unsafe_allow_html=True)
        share = st.slider("RFN-Anteil am Sponsoringwert (%)", 0, 100, int(s["share"]))
        prov = st.slider("Vertriebsprovision (% vom RFN-Anteil)", 0, 100, int(s["provision"]))
        ex = 1000; rf = ex * share / 100; pv = rf * prov / 100
        st.markdown(f'<p class="muted">Beispiel 1.000 €: Verein {eur(ex - rf)} · RFN netto {eur(rf - pv)} · Provision {eur(pv)}</p>', unsafe_allow_html=True)
        st.markdown('<div class="sec">Team / Vertriebler</div>', unsafe_allow_html=True)
        team = st.text_area("Ein Name pro Zeile", "\n".join(s["team"]), height=110, label_visibility="collapsed")
        st.markdown('<div class="sec">Kostenbremse Google</div>', unsafe_allow_html=True)
        gmax = st.number_input("Max. Google-Abfragen pro Recherche", 100, 20000, int(s.get("google_max", 1000)), step=100,
                               help="Jede Abfrage ist eine kostenpflichtige Google-Anfrage (nach dem monatlichen Freikontingent). "
                                    "Kleinstadt ≈ 150–400 Abfragen, Großstadt deutlich mehr. Ist die Grenze erreicht, "
                                    "werden dichte Gebiete nicht weiter zerlegt – alle Suchbegriffe laufen trotzdem einmal.")
        st.markdown('<p class="muted">Aktuelle Preise: mapsplatform.google.com/pricing · Verbrauch: Google Cloud → Abrechnung → Berichte</p>',
                    unsafe_allow_html=True)
        if st.button("💾  Speichern", type="primary"):
            put("settings", {"share": share, "provision": prov, "google_max": int(gmax),
                             "team": [t.strip() for t in team.splitlines() if t.strip()]})
            st.toast("Einstellungen gespeichert", icon="✅"); st.rerun()
    with r.container(border=True):
        st.markdown('<div class="sec">Datenspeicher</div>', unsafe_allow_html=True)
        if storage.is_remote() and not st.session_state.get("db_error"):
            st.markdown('<div class="ok">✓ Supabase verbunden – alle Daten werden dauerhaft gespeichert.</div>', unsafe_allow_html=True)
        elif st.session_state.get("db_error"):
            st.markdown(f'<div class="warn">⚠ Datenbankfehler: {st.session_state.db_error}</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="warn">⚠ Lokaler Speicher. Auf Streamlit Cloud gehen Daten bei einem Neustart verloren. '
                        'Supabase verbinden (Anleitung SETUP.md) oder regelmäßig Backup ziehen.</div>', unsafe_allow_html=True)
        st.write("")
        st.download_button("⬇️  Komplett-Backup (JSON)", json.dumps(DB(), ensure_ascii=False, indent=1).encode("utf-8"),
                           file_name=f"RFN_Cockpit_Backup_{today()}.json", width="stretch")
        up = st.file_uploader("Backup einspielen", type="json")
        if up is not None and st.button("Backup wiederherstellen"):
            data = json.load(up)
            for k, v in data.items(): put(k, v)
            st.toast("Backup eingespielt", icon="✅"); st.rerun()
        st.markdown("---")
        st.markdown('<div class="sec">Vereine verwalten</div>', unsafe_allow_html=True)
        C = clubs()
        if C:
            cid = st.selectbox("Verein", [c["id"] for c in C], format_func=lambda x: club_by_id(x)["name"], key="set_club")
            club = club_by_id(cid)
            nm = st.text_input("Name", club["name"], key=f"nm_{cid}")
            ad = st.text_input("Standort", club.get("adresse", ""), key=f"ad_{cid}")
            a, b = st.columns(2)
            if a.button("Übernehmen", key=f"ren_{cid}"):
                for c in C:
                    if c["id"] == cid: c.update(name=nm.strip() or c["name"], adresse=ad.strip())
                put("clubs", C); st.rerun()
            sure = b.checkbox(f"Löschen bestätigen ({len(leads(cid))} Unternehmen)", key=f"del_{cid}")
            if sure and b.button("🗑️ Verein löschen", key=f"delb_{cid}"):
                put("clubs", [c for c in C if c["id"] != cid])
                try: storage.delete(f"leads:{cid}")
                except Exception: pass
                DB().pop(f"leads:{cid}", None); st.session_state.club = "__all__"; st.rerun()
        else:
            st.markdown('<p class="muted">Noch keine Vereine.</p>', unsafe_allow_html=True)


# ══════════════════════════════ Ablauf ══════════════════════════════
DB()
if "club" not in st.session_state: st.session_state.club = "__all__"
if "nav" not in st.session_state: st.session_state.nav = PAGES[0]
if "pending_start" in st.session_state:
    start_job(st.session_state.pop("pending_start"))
if "goto_club" in st.session_state: st.session_state.club = st.session_state.pop("goto_club")
if "goto_nav" in st.session_state: st.session_state.nav = st.session_state.pop("goto_nav")

open_dialog = False
with st.sidebar:
    st.markdown(BRAND, unsafe_allow_html=True)
    if st.button("＋  Neue Recherche", type="primary", width="stretch"):
        open_dialog = True
    st.write("")
    st.radio("Navigation", PAGES, key="nav", label_visibility="collapsed")
    st.markdown("---")
    copts = ["__all__"] + [c["id"] for c in clubs()]
    if st.session_state.club not in copts: st.session_state.club = "__all__"
    st.selectbox("Verein", copts, key="club",
                 format_func=lambda x: "Alle Vereine" if x == "__all__" else club_by_id(x)["name"])
    st.markdown("---")
    st.caption("🟢 Datenbank verbunden" if storage.is_remote() and not st.session_state.get("db_error") else "🟠 Lokaler Speicher – Backup ziehen")
    st.caption(f"Quelle: {'Google + OpenStreetMap' if secret('GOOGLE_API_KEY', '') else 'OpenStreetMap'}")
    st.caption(f"Engine: {getattr(rp, 'VERSION', 'alt – bitte Reboot')}")
    if pw_required and st.button("Abmelden", width="stretch"):
        st.session_state.auth = False; st.rerun()

if open_dialog:
    research_dialog()

if st.session_state.get("db_error"):
    st.markdown(f'<div class="warn">⚠ Datenbank nicht erreichbar: {st.session_state.db_error}</div>', unsafe_allow_html=True)
elif not storage.is_remote():
    st.markdown('<div class="warn">⚠ Daten liegen im lokalen Speicher und gehen auf Streamlit Cloud bei einem Neustart verloren. '
                'Supabase verbinden (SETUP.md) – bis dahin unter ⚙️ Einstellungen regelmäßig ein Backup ziehen.</div>', unsafe_allow_html=True)
if st.session_state.get("flash"):
    st.success(st.session_state.pop("flash"))
if st.session_state.get("flash_error"):
    st.error(st.session_state.pop("flash_error"))
if any(j["status"] == "running" or not j.get("shown") for j in REG["jobs"].values()):
    job_panel()

{PAGES[0]: page_dashboard, PAGES[1]: page_leads, PAGES[2]: page_sponsoren, PAGES[3]: page_leistungen,
 PAGES[4]: page_runs, PAGES[5]: page_settings}[st.session_state.nav]()
